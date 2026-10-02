import os

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login as auth_login, logout as auth_logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.exceptions import PermissionDenied
from django.http import FileResponse, Http404, HttpResponse
from django.db.models import Q, Count
from django.utils.dateparse import parse_date
from django.utils import timezone as tz
from django.core.paginator import Paginator
from django.db import IntegrityError
from django.contrib.postgres.search import SearchVector, SearchQuery

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from urllib.parse import quote

from .models import (
    Document, Comment, Tag, AuditLog, Folder, Directorate, Department,
    CustomUser, DocumentAttachment, DocumentType, DocumentVersion,
)
from .permissions import (
    get_visible_documents, can_view_document, can_edit_document, can_delete_document,
    is_supervisory_action, get_visible_folders, can_view_folder, can_manage_folder,
    can_delete_folder,
)
from .audit import log_action
from .utils.files import sanitize_folder_name


# ==========================================================
# تسجيل الدخول / الخروج
# ==========================================================

def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')

        user = authenticate(request, username=username, password=password)

        if user is not None:
            if user.is_system_admin and not user.role:
                messages.error(request, 'هذا الحساب مخصص لإدارة النظام، الرجاء استخدام لوحة الإدارة.')
                return render(request, 'core/login.html')

            auth_login(request, user)
            return redirect('dashboard')
        else:
            messages.error(request, 'اسم المستخدم أو كلمة المرور غير صحيحة.')

    return render(request, 'core/login.html')


@login_required
def logout_view(request):
    auth_logout(request)
    return redirect('login')


# ==========================================================
# لوحة التحكم
# ==========================================================

@login_required
def dashboard(request):
    documents = get_visible_documents(request.user)

    context = {
        'total_documents': documents.count(),
        'internal_count': documents.filter(source=Document.Source.INTERNAL).count(),
        'external_count': documents.filter(source=Document.Source.EXTERNAL).count(),
    }
    return render(request, 'core/dashboard.html', context)


# ==========================================================
# المجلدات
# ==========================================================

@login_required
def folder_list(request):
    folders_list = get_visible_folders(request.user).filter(parent__isnull=True).select_related('created_by').annotate(
        doc_count=Count('documents', distinct=True),
        subfolder_count=Count('subfolders', distinct=True),
    ).order_by('name')
    documents_list = get_visible_documents(request.user).filter(folder__isnull=True).select_related('department', 'directorate', 'uploaded_by')

    folders_paginator = Paginator(folders_list, 25)
    folders_page = folders_paginator.get_page(request.GET.get('folders_page'))

    documents_paginator = Paginator(documents_list, 25)
    documents_page = documents_paginator.get_page(request.GET.get('page'))

    context = {
        'folders': folders_page,
        'page_obj': documents_page,
        'current_folder': None,
    }
    return render(request, 'core/folder_browser.html', context)


@login_required
def folder_detail(request, folder_id):
    folder = get_object_or_404(Folder, id=folder_id)

    if not can_view_folder(request.user, folder):
        raise PermissionDenied

    subfolders_list = get_visible_folders(request.user).filter(parent_id=folder.id).select_related('created_by').annotate(
        doc_count=Count('documents', distinct=True),
        subfolder_count=Count('subfolders', distinct=True),
    ).order_by('name')
    documents_list = get_visible_documents(request.user).filter(folder_id=folder.id).select_related('department', 'directorate', 'uploaded_by')

    folders_paginator = Paginator(subfolders_list, 25)
    folders_page = folders_paginator.get_page(request.GET.get('folders_page'))

    documents_paginator = Paginator(documents_list, 25)
    documents_page = documents_paginator.get_page(request.GET.get('page'))

    context = {
        'folders': folders_page,
        'page_obj': documents_page,
        'current_folder': folder,
        'can_delete': can_delete_folder(request.user, folder),
    }
    return render(request, 'core/folder_browser.html', context)


@login_required
def create_folder(request):
    if not request.user.role:
        raise PermissionDenied

    if not request.user.directorate_id:
        messages.error(request, 'حسابك غير مرتبط بمديرية، الرجاء مراجعة الإدارة.')
        return redirect('dashboard')

    department = None
    if request.user.role in [CustomUser.Role.EMPLOYEE, CustomUser.Role.HEAD_OF_DEPT]:
        department = request.user.department
        if not department:
            messages.error(request, 'حسابك غير مرتبط بدائرة، الرجاء مراجعة الإدارة.')
            return redirect('dashboard')
    elif request.user.role == CustomUser.Role.DIRECTORATE_MANAGER:
        department_id = request.POST.get('department_id') or request.GET.get('department_id')
        if department_id:
            try:
                department = Department.objects.get(
                    id=department_id,
                    directorate_id=request.user.directorate_id
                )
            except Department.DoesNotExist:
                raise PermissionDenied

    parent_id = request.POST.get('parent_id') or request.GET.get('parent_id')
    parent = None
    if parent_id:
        parent = get_object_or_404(Folder, id=parent_id)
        if not can_view_folder(request.user, parent):
            raise PermissionDenied

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        if not name:
            messages.error(request, 'الرجاء إدخال اسم المجلد.')
        else:
            if request.user.role == CustomUser.Role.DIRECTORATE_MANAGER and not department:
                messages.error(request, 'الرجاء اختيار الدائرة.')
                departments = Department.objects.filter(directorate_id=request.user.directorate_id)
                return render(request, 'core/create_folder.html', {
                    'parent': parent,
                    'departments': departments,
                })

            folder = Folder.objects.create(
                name=name,
                parent=parent,
                directorate=request.user.directorate,
                department=department,
                created_by=request.user,
            )
            messages.success(request, f'تم إنشاء المجلد "{folder.name}" بنجاح.')
            if parent:
                return redirect('folder_detail', folder_id=parent.id)
            return redirect('folder_list')

    departments = None
    if request.user.role == CustomUser.Role.DIRECTORATE_MANAGER:
        departments = Department.objects.filter(directorate_id=request.user.directorate_id)

    return render(request, 'core/create_folder.html', {
        'parent': parent,
        'departments': departments,
    })


@login_required
def delete_folder(request, folder_id):
    folder = get_object_or_404(Folder, id=folder_id)

    if not can_delete_folder(request.user, folder):
        raise PermissionDenied

    if folder.has_content():
        messages.error(request, 'لا يمكن حذف المجلد لأنه يحتوي على وثائق أو مجلدات فرعية. الرجاء تفريغه أولاً.')
        return redirect('folder_detail', folder_id=folder.id)

    if request.method == 'POST':
        parent_id = folder.parent_id
        folder.delete()
        messages.success(request, 'تم حذف المجلد.')
        if parent_id:
            return redirect('folder_detail', folder_id=parent_id)
        return redirect('folder_list')

    return render(request, 'core/confirm_delete_folder.html', {'folder': folder})


# ==========================================================
# رفع وثيقة
# ==========================================================

@login_required
def upload_document(request):
    if not request.user.role:
        raise PermissionDenied

    if request.user.role == CustomUser.Role.GENERAL_MANAGER:
        raise PermissionDenied

    if not request.user.directorate_id:
        messages.error(request, 'حسابك غير مرتبط بمديرية، الرجاء مراجعة الإدارة.')
        return redirect('dashboard')

    department = None
    if request.user.role in [CustomUser.Role.EMPLOYEE, CustomUser.Role.HEAD_OF_DEPT]:
        department = request.user.department
        if not department:
            messages.error(request, 'حسابك غير مرتبط بدائرة، الرجاء مراجعة الإدارة.')
            return redirect('dashboard')
    elif request.user.role == CustomUser.Role.DIRECTORATE_MANAGER:
        department_id = request.POST.get('department_id') or request.GET.get('department_id')
        if department_id:
            try:
                department = Department.objects.get(
                    id=department_id,
                    directorate_id=request.user.directorate_id
                )
            except Department.DoesNotExist:
                raise PermissionDenied

    folder_id = request.POST.get('folder_id') or request.GET.get('folder_id')
    folder = None
    if folder_id:
        folder = get_object_or_404(Folder, id=folder_id)
        if not can_view_folder(request.user, folder):
            raise PermissionDenied

    if request.method == 'POST':
        document_number = request.POST.get('document_number', '').strip()
        title = request.POST.get('title', '').strip()
        document_date = request.POST.get('document_date')
        export_date = request.POST.get('export_date', '').strip()
        document_type_id = request.POST.get('document_type_id', '').strip()
        source = request.POST.get('source')
        uploaded_file = request.FILES.get('file')
        tags_raw = request.POST.get('tags', '')

        issuing_directorate_id = request.POST.get('issuing_directorate')
        external_entity_name = request.POST.get('external_entity_name', '').strip()
        destination_entity_name = request.POST.get('destination_entity_name', '').strip()
        diwan_number = request.POST.get('diwan_number', '').strip()

        status = request.POST.get('status', '').strip()

        if not (document_number and title and document_date and source and uploaded_file):
            return _upload_error(request, folder, 'الرجاء تعبئة كل الحقول المطلوبة واختيار ملف.')

        if not document_number[0].isdigit():
            return _upload_error(request, folder, 'رقم الوثيقة يجب أن يبدأ برقم.')

        if source not in Document.Source.values:
            return _upload_error(request, folder, 'قيمة المصدر غير صحيحة.')

        if request.user.role == CustomUser.Role.DIRECTORATE_MANAGER and not department:
            return _upload_error(request, folder, 'الرجاء اختيار الدائرة.')

        parsed_date = parse_date(document_date)
        if not parsed_date:
            return _upload_error(request, folder, 'تاريخ الوثيقة غير صالحة.')
        document_year = parsed_date.year

        parsed_export_date = parse_date(export_date) if export_date else None

        document_type = None
        if document_type_id:
            try:
                document_type = DocumentType.objects.get(id=document_type_id)
            except DocumentType.DoesNotExist:
                document_type = None

        duplicate_exists = Document.objects.filter(
            directorate=request.user.directorate,
            document_year=document_year,
            document_number=document_number
        ).exists()
        if duplicate_exists:
            return _upload_error(
                request, folder,
                f'رقم الوثيقة "{document_number}" مستخدم مسبقاً ضمن مديرية {request.user.directorate.name} لسنة {document_year}. الرجاء اختيار رقم آخر.'
            )

        try:
            doc = Document.objects.create(
                document_number=document_number,
                title=title,
                file=uploaded_file,
                document_date=parsed_date,
                export_date=parsed_export_date,
                document_type=document_type,
                source=source,
                issuing_directorate_id=issuing_directorate_id if source == Document.Source.INTERNAL else None,
                external_entity_name=external_entity_name if source == Document.Source.EXTERNAL else '',
                destination_entity_name=destination_entity_name or None,
                diwan_number=diwan_number or None,
                status=status or None,
                folder=folder,
                directorate=request.user.directorate,
                department=department,
                uploaded_by=request.user,
            )
        except IntegrityError:
            return _upload_error(
                request, folder,
                f'رقم الوثيقة "{document_number}" مستخدم مسبقاً ضمن مديرية {request.user.directorate.name} لسنة {document_year}. الرجاء اختيار رقم آخر.'
            )

        for tag_name in [t.strip() for t in tags_raw.split(',') if t.strip()]:
            tag, _ = Tag.objects.get_or_create(name=tag_name)
            doc.tags.add(tag)

        log_action(request.user, doc, AuditLog.Action.CREATE)

        messages.success(request, f'تم رفع الوثيقة رقم {doc.document_number} بنجاح.')
        return redirect('document_detail', doc_id=doc.id)

    departments = None
    if request.user.role == CustomUser.Role.DIRECTORATE_MANAGER:
        departments = Department.objects.filter(directorate_id=request.user.directorate_id)
    return render(request, 'core/upload.html', _upload_context(folder, departments=departments))


def _upload_context(folder, departments=None):
    return {
        'folder': folder,
        'directorates': Directorate.objects.all(),
        'departments': departments,
        'document_types': DocumentType.objects.all(),
    }


def _upload_error(request, folder, message, departments=None):
    if departments is None and request.user.role == CustomUser.Role.DIRECTORATE_MANAGER:
        departments = Department.objects.filter(directorate_id=request.user.directorate_id)
    messages.error(request, message)
    return render(request, 'core/upload.html', _upload_context(folder, departments=departments))


# ==========================================================
# البحث
# ==========================================================

def _apply_search_filters(documents, request):
    """
    تُطبّق كل فلاتر البحث على queryset — تُستدعى من:
    - search_documents
    - export_documents_excel
    لتضمن أن الاثنين يعطيان نفس النتائج دائماً.
    """
    document_number = request.GET.get('document_number', '').strip()
    name = request.GET.get('name', '').strip()
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    export_date_from = request.GET.get('export_date_from', '').strip()
    export_date_to = request.GET.get('export_date_to', '').strip()
    document_type_id = request.GET.get('document_type_id', '').strip()
    destination_entity_name = request.GET.get('destination_entity_name', '').strip()
    diwan_number = request.GET.get('diwan_number', '').strip()
    keyword = request.GET.get('keyword', '').strip()
    source = request.GET.get('source', '').strip()
    issuing_directorate_id = request.GET.get('issuing_directorate', '').strip()
    external_entity_name = request.GET.get('external_entity_name', '').strip()
    status = request.GET.get('status', '').strip()

    if document_number:
        documents = documents.filter(document_number__icontains=document_number)

    if name:
        documents = documents.filter(title__icontains=name)

    if date_from:
        documents = documents.filter(document_date__gte=date_from)
    if date_to:
        documents = documents.filter(document_date__lte=date_to)

    if export_date_from:
        documents = documents.filter(export_date__gte=export_date_from)
    if export_date_to:
        documents = documents.filter(export_date__lte=export_date_to)

    if document_type_id:
        documents = documents.filter(document_type_id=document_type_id)

    if destination_entity_name:
        documents = documents.filter(destination_entity_name__icontains=destination_entity_name)

    if diwan_number:
        documents = documents.filter(diwan_number__icontains=diwan_number)

    if status:
        documents = documents.filter(status__icontains=status)

    if keyword:
        documents = documents.annotate(
            search=SearchVector(
                'title', 'document_number', 'external_entity_name',
                'tags__name', 'diwan_number', 'destination_entity_name',
                'status',
            )
        ).filter(
            Q(search=SearchQuery(keyword)) |
            Q(tags__name__icontains=keyword) |
            Q(document_number__icontains=keyword) |
            Q(diwan_number__icontains=keyword)
        ).distinct()

    if source in Document.Source.values:
        documents = documents.filter(source=source)

    if source == Document.Source.INTERNAL and issuing_directorate_id:
        documents = documents.filter(issuing_directorate_id=issuing_directorate_id)
    elif source == Document.Source.EXTERNAL and external_entity_name:
        documents = documents.filter(external_entity_name__icontains=external_entity_name)

    return documents


@login_required
def search_documents(request):
    # نقرأ الفلاتر أولاً لمعرفة إن كان هناك بحث
    document_number = request.GET.get('document_number', '').strip()
    name = request.GET.get('name', '').strip()
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    export_date_from = request.GET.get('export_date_from', '').strip()
    export_date_to = request.GET.get('export_date_to', '').strip()
    document_type_id = request.GET.get('document_type_id', '').strip()
    destination_entity_name = request.GET.get('destination_entity_name', '').strip()
    diwan_number = request.GET.get('diwan_number', '').strip()
    keyword = request.GET.get('keyword', '').strip()
    source = request.GET.get('source', '').strip()
    issuing_directorate_id = request.GET.get('issuing_directorate', '').strip()
    external_entity_name = request.GET.get('external_entity_name', '').strip()
    status = request.GET.get('status', '').strip()

    has_search_params = any([
        document_number, name, date_from, date_to,
        export_date_from, export_date_to, document_type_id,
        destination_entity_name, diwan_number, keyword,
        source, issuing_directorate_id, external_entity_name, status,
    ])

    if not has_search_params:
        context = {
            'documents': None,
            'page_obj': None,
            'has_searched': False,
            'directorates': Directorate.objects.all(),
            'document_types': DocumentType.objects.all(),
            'query_string': '',
        }
        return render(request, 'core/search.html', context)

    # ✅ استدعاء واحد لكل الفلاتر
    documents = get_visible_documents(request.user).select_related(
        'department', 'directorate', 'uploaded_by', 'folder', 'document_type'
    )
    documents = _apply_search_filters(documents, request)
    documents = documents.order_by('-document_date', '-id')

    paginator = Paginator(documents, 25)
    page_obj = paginator.get_page(request.GET.get('page'))

    query_params = request.GET.copy()
    query_params.pop('page', None)
    query_string = query_params.urlencode()

    context = {
        'documents': page_obj,
        'page_obj': page_obj,
        'has_searched': True,
        'directorates': Directorate.objects.all(),
        'document_types': DocumentType.objects.all(),
        'query_string': query_string,
    }
    return render(request, 'core/search.html', context)


# ==========================================================
# ✅✅✅  تصدير Excel (شامل المجلدات الفرعية)  ✅✅✅
# ==========================================================

def _get_all_folder_ids_recursive(folder_id):
    """
    ترجع قائمة بكل معرفات المجلدات: المجلد نفسه + كل المجلدات الفرعية
    (بشكل تعاقبي، أي الجيل الأول والثاني والثالث... إلى ما لا نهاية).
    """
    ids = [folder_id]
    queue = [folder_id]
    while queue:
        current = queue.pop()
        child_ids = list(
            Folder.objects.filter(parent_id=current).values_list('id', flat=True)
        )
        for cid in child_ids:
            ids.append(cid)
            queue.append(cid)
    return ids


# ✅ حد الإشعار: عند تجاوز هذا العدد، يظهر تحذير للمستخدم قبل التصدير
EXPORT_WARNING_THRESHOLD = 10000


@login_required
def export_documents_excel(request):
    """
    تصدير نتائج البحث إلى ملف Excel (xlsx).
    - يستخدم نفس فلاتر البحث في search_documents تماماً.
    - عند تجاوز 10,000 وثيقة → يظهر تحذير للمستخدم قبل التنفيذ.
    - يقبل معاملاً إضافياً `folder_id` لتقييد التصدير بمجلد معين.
    - عند تمرير `folder_id` → يشمل المجلد + كل المجلدات الفرعية تحته.
    - عند تمرير `folder_id=none` → فقط الوثائق غير المصنفة.
    - لا يحتوي على الملف الأصلي ولا المرفقات — فقط بيانات وصفية.
    """
    # ✅ فلتر المجلد (خاص بالتصدير فقط، لا علاقة له بفلاتر البحث)
    folder_id = request.GET.get('folder_id', '').strip()

    # ✅ queryset أساسي
    documents = get_visible_documents(request.user).select_related(
        'department', 'directorate', 'uploaded_by', 'folder', 'document_type',
        'issuing_directorate',
    ).prefetch_related('tags')

    # ✅ فلتر المجلد أولاً (شامل المجلدات الفرعية)
    if folder_id:
        if folder_id == 'none':
            documents = documents.filter(folder__isnull=True)
        else:
            try:
                fid = int(folder_id)
                folder_ids = _get_all_folder_ids_recursive(fid)
                documents = documents.filter(folder_id__in=folder_ids)
            except (ValueError, TypeError):
                pass

    # ✅ ثم كل فلاتر البحث المشتركة (نفس الدالة المستخدمة في الشاشة)
    documents = _apply_search_filters(documents, request)
    documents = documents.order_by('-document_date', '-id')

    # ✅✅✅ حماية: إشعار عند تجاوز الحد الأقصى ✅✅✅
    total_count = documents.count()

    if total_count > EXPORT_WARNING_THRESHOLD and request.GET.get('confirmed') != '1':
        context = {
            'total_count': total_count,
            'threshold': EXPORT_WARNING_THRESHOLD,
            'query_string': request.GET.urlencode(),
        }
        return render(request, 'core/confirm_export.html', context)

    # ✅✅✅ كل بناء الملف داخل try لضمان التعامل مع أي خطأ ✅✅✅
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'الوثائق'
        ws.sheet_view.rightToLeft = True

        headers = [
            'رقم الوثيقة',
            'السنة',
            'الرقم الكامل (رقم/سنة)',
            'اسم الوثيقة',
            'نوع الوثيقة',
            'تاريخ الوثيقة',
            'تاريخ التصدير/الختم',
            'رقم الديوان',
            'المصدر',
            'المديرية الصادرة عنها',
            'الجهة الخارجية',
            'الجهة المحولة إليها',
            'الوسوم',
            'حالة الوثيقة',
            'المديرية',
            'الدائرة',
            'المجلد',
            'رفعها',
            'تاريخ الرفع',
        ]

        header_font = Font(bold=True, color='FFFFFF', size=12)
        header_fill = PatternFill(start_color='0F4A42', end_color='0F4A42', fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        data_align = Alignment(horizontal='right', vertical='center', wrap_text=True)
        thin = Side(border_style='thin', color='D0D0D0')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_align
            c.border = border

        for row_idx, doc in enumerate(documents, start=2):
            tags_str = ', '.join(t.name for t in doc.tags.all())
            row = [
                doc.document_number,
                doc.document_year,
                doc.full_number,
                doc.title,
                doc.document_type.name if doc.document_type else '',
                doc.document_date.strftime('%Y-%m-%d') if doc.document_date else '',
                doc.export_date.strftime('%Y-%m-%d') if doc.export_date else '',
                doc.diwan_number or '',
                doc.get_source_display(),
                doc.issuing_directorate.name if doc.issuing_directorate else '',
                doc.external_entity_name or '',
                doc.destination_entity_name or '',
                tags_str,
                doc.status or '',
                doc.directorate.name if doc.directorate else '',
                doc.department.name if doc.department else '',
                doc.folder.name if doc.folder else 'غير مصنف',
                str(doc.uploaded_by) if doc.uploaded_by else '',
                doc.created_at.strftime('%Y-%m-%d %H:%M') if doc.created_at else '',
            ]
            for col_idx, val in enumerate(row, start=1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.alignment = data_align
                c.border = border

        column_widths = [14, 8, 18, 35, 16, 14, 16, 12, 10, 22, 22, 22, 25, 18, 18, 18, 20, 22, 18]
        for i, w in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        ws.freeze_panes = 'A2'

        # اسم الملف
        folder_part = ''
        if folder_id and folder_id != 'none':
            try:
                f_obj = Folder.objects.filter(id=int(folder_id)).first()
                if f_obj:
                    folder_part = f"_{sanitize_folder_name(f_obj.name)}"
            except (ValueError, TypeError):
                pass
        elif folder_id == 'none':
            folder_part = '_uncategorized'

        filename = f"documents{folder_part}_{tz.now().strftime('%Y-%m-%d_%H-%M')}.xlsx"
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(filename)}"

        wb.save(response)
        return response

    except MemoryError:
        messages.error(request, 'الملف كبير جداً. صغّر نطاق البحث ثم أعد المحاولة.')
        return redirect('search')
    except Exception as e:
        messages.error(request, f'تعذّر إنشاء ملف Excel: {str(e)[:200]}')
        return redirect('search')


# ==========================================================
# عرض تفاصيل وثيقة / تحميل / معاينة / تعديل / حذف / مرفقات
# ==========================================================

@login_required
def document_detail(request, doc_id):
    document = get_object_or_404(Document, id=doc_id)

    if not can_view_document(request.user, document):
        raise PermissionDenied

    context = {
        'document': document,
        'can_edit': can_edit_document(request.user, document),
        'can_delete': can_delete_document(request.user, document),
        'comments': document.comments.select_related('created_by'),
        'attachments': document.attachments.select_related('uploaded_by'),
        'versions': document.versions.select_related('uploaded_by'),
    }
    return render(request, 'core/document_detail.html', context)


@login_required
def download_document(request, doc_id):
    document = get_object_or_404(Document, id=doc_id)

    if not can_view_document(request.user, document):
        raise PermissionDenied

    log_action(request.user, document, AuditLog.Action.DOWNLOAD)

    try:
        return FileResponse(document.file.open('rb'), as_attachment=True,
                             filename=document.file.name.split('/')[-1])
    except FileNotFoundError:
        raise Http404('الملف غير موجود على الخادم.')


@login_required
def preview_document(request, doc_id):
    document = get_object_or_404(Document, id=doc_id)

    if not can_view_document(request.user, document):
        raise PermissionDenied

    try:
        import mimetypes
        mime_type, _ = mimetypes.guess_type(document.file.name)
        if mime_type is None:
            mime_type = 'application/octet-stream'

        response = FileResponse(document.file.open('rb'), content_type=mime_type)
        response['Content-Disposition'] = f'inline; filename="{document.file.name.split("/")[-1]}"'
        return response
    except FileNotFoundError:
        raise Http404('الملف غير موجود على الخادم.')


@login_required
def upload_attachment(request, doc_id):
    document = get_object_or_404(Document, id=doc_id)

    if not can_edit_document(request.user, document):
        raise PermissionDenied

    if request.method == 'POST' and request.FILES.get('file'):
        DocumentAttachment.objects.create(
            document=document,
            file=request.FILES['file'],
            uploaded_by=request.user
        )
        messages.success(request, 'تمت إضافة الملف المرفق بنجاح.')
    else:
        messages.error(request, 'الرجاء اختيار ملف أولاً.')

    return redirect('document_detail', doc_id=document.id)


@login_required
def delete_attachment(request, attachment_id):
    attachment = get_object_or_404(DocumentAttachment, id=attachment_id)
    document = attachment.document

    if not can_view_document(request.user, document):
        raise PermissionDenied

    if not can_delete_document(request.user, document):
        raise PermissionDenied

    if request.method == 'POST':
        attachment.delete()
        messages.success(request, 'تم حذف الملف المرفق بنجاح.')
        return redirect('document_detail', doc_id=document.id)

    return redirect('document_detail', doc_id=document.id)


@login_required
def replace_attachment(request, attachment_id):
    attachment = get_object_or_404(DocumentAttachment, id=attachment_id)
    document = attachment.document

    if not can_view_document(request.user, document):
        raise PermissionDenied

    if not can_edit_document(request.user, document):
        raise PermissionDenied

    if request.method == 'POST' and request.FILES.get('file'):
        attachment.file = request.FILES['file']
        attachment.uploaded_by = request.user
        attachment.save()
        messages.success(request, 'تم استبدال الملف المرفق بنجاح.')
    else:
        messages.error(request, 'الرجاء اختيار ملف أولاً.')

    return redirect('document_detail', doc_id=document.id)


@login_required
def update_document_file(request, doc_id):
    document = get_object_or_404(Document, id=doc_id)

    if not can_edit_document(request.user, document):
        raise PermissionDenied

    if request.method == 'POST':
        new_file = request.FILES.get('file')
        note = request.POST.get('note', '').strip()

        if not new_file:
            messages.error(request, 'الرجاء اختيار ملف جديد.')
            return render(request, 'core/update_document_file.html', {'document': document})

        if document.file:
            DocumentVersion.objects.create(
                document=document,
                file_name=document.file.name,
                original_filename=os.path.basename(document.file.name),
                note=note,
                uploaded_by=request.user,
            )

        document.file = new_file
        document.save()

        log_action(request.user, document, AuditLog.Action.EDIT)
        messages.success(request, 'تم تحديث ملف الوثيقة بنجاح. النسخة السابقة محفوظة.')
        return redirect('document_detail', doc_id=document.id)

    return render(request, 'core/update_document_file.html', {'document': document})


# ==========================================================
# تعديل وثيقة
# ==========================================================

@login_required
def edit_document(request, doc_id):
    document = get_object_or_404(Document, id=doc_id)

    if not can_edit_document(request.user, document):
        raise PermissionDenied

    if request.method == 'POST':
        title = request.POST.get('title', '').strip()
        document_number = request.POST.get('document_number', '').strip()
        document_date = request.POST.get('document_date')
        export_date = request.POST.get('export_date', '').strip()
        document_type_id = request.POST.get('document_type_id', '').strip()
        source = request.POST.get('source', '').strip()
        issuing_directorate_id = request.POST.get('issuing_directorate', '').strip()
        external_entity_name = request.POST.get('external_entity_name', '').strip()
        destination_entity_name = request.POST.get('destination_entity_name', '').strip()
        diwan_number = request.POST.get('diwan_number', '').strip()
        tags_raw = request.POST.get('tags', '')
        status = request.POST.get('status', '').strip()

        if not (title and document_number and document_date and source):
            messages.error(request, 'الرجاء تعبئة جميع الحقول المطلوبة.')
            return render(request, 'core/edit_document.html', _edit_context(document))

        if not document_number[0].isdigit():
            messages.error(request, 'رقم الوثيقة يجب أن يبدأ برقم.')
            return render(request, 'core/edit_document.html', _edit_context(document))

        if source not in Document.Source.values:
            messages.error(request, 'قيمة المصدر غير صحيحة.')
            return render(request, 'core/edit_document.html', _edit_context(document))

        parsed_date = parse_date(document_date)
        if not parsed_date:
            messages.error(request, 'تاريخ غير صالح.')
            return render(request, 'core/edit_document.html', _edit_context(document))

        duplicate_exists = Document.objects.filter(
            directorate=document.directorate,
            document_year=parsed_date.year,
            document_number=document_number
        ).exclude(pk=document.pk).exists()

        if duplicate_exists:
            messages.error(
                request,
                f'رقم الوثيقة "{document_number}" مستخدم مسبقاً ضمن مديرية {document.directorate.name} لسنة {parsed_date.year}. الرجاء اختيار رقم آخر.'
            )
            return render(request, 'core/edit_document.html', _edit_context(document))

        document.title = title
        document.document_number = document_number
        document.document_date = parsed_date
        document.export_date = parse_date(export_date) if export_date else None
        document.diwan_number = diwan_number or None
        document.source = source
        document.status = status or None

        if document_type_id:
            try:
                document.document_type = DocumentType.objects.get(id=document_type_id)
            except DocumentType.DoesNotExist:
                document.document_type = None
        else:
            document.document_type = None

        if source == Document.Source.INTERNAL and issuing_directorate_id:
            try:
                document.issuing_directorate = Directorate.objects.get(id=issuing_directorate_id)
            except Directorate.DoesNotExist:
                document.issuing_directorate = None
            document.external_entity_name = None
        elif source == Document.Source.EXTERNAL:
            document.external_entity_name = external_entity_name or None
            document.issuing_directorate = None
        else:
            document.issuing_directorate = None
            document.external_entity_name = None

        document.destination_entity_name = destination_entity_name or None

        try:
            document.save()
        except IntegrityError:
            messages.error(
                request,
                f'رقم الوثيقة "{document_number}" مستخدم مسبقاً ضمن مديرية {document.directorate.name} لسنة {parsed_date.year}. الرجاء اختيار رقم آخر.'
            )
            return render(request, 'core/edit_document.html', _edit_context(document))

        document.tags.clear()
        for tag_name in [t.strip() for t in tags_raw.split(',') if t.strip()]:
            tag, _ = Tag.objects.get_or_create(name=tag_name)
            document.tags.add(tag)

        log_action(request.user, document, AuditLog.Action.EDIT)
        messages.success(request, 'تم تعديل الوثيقة بنجاح.')
        return redirect('document_detail', doc_id=document.id)

    return render(request, 'core/edit_document.html', _edit_context(document))


def _edit_context(document):
    return {
        'document': document,
        'document_types': DocumentType.objects.all(),
        'directorates': Directorate.objects.all(),
    }


@login_required
def delete_document(request, doc_id):
    document = get_object_or_404(Document, id=doc_id)

    if not can_delete_document(request.user, document):
        raise PermissionDenied

    if request.method == 'POST':
        log_action(request.user, document, AuditLog.Action.DELETE)
        document.delete()
        messages.success(request, 'تم حذف الوثيقة، وتم تسجيل الإجراء بسجل التدقيق.')
        return redirect('dashboard')

    return render(request, 'core/confirm_delete.html', {'document': document})


@login_required
def add_comment(request, doc_id):
    document = get_object_or_404(Document, id=doc_id)

    if not can_view_document(request.user, document):
        raise PermissionDenied

    if request.method == 'POST':
        text = request.POST.get('text', '').strip()
        if text:
            Comment.objects.create(document=document, text=text, created_by=request.user)

    return redirect('document_detail', doc_id=document.id)


# ==========================================================
# تغيير كلمة المرور
# ==========================================================

@login_required
def change_password(request):
    from django.contrib.auth.password_validation import validate_password
    from django.core.exceptions import ValidationError

    if request.method == 'POST':
        old_password = request.POST.get('old_password', '')
        new_password = request.POST.get('new_password', '')
        confirm_password = request.POST.get('confirm_password', '')

        if not request.user.check_password(old_password):
            messages.error(request, 'كلمة المرور الحالية غير صحيحة.')
            return render(request, 'core/change_password.html')

        if new_password != confirm_password:
            messages.error(request, 'كلمتا المرور الجديدتان غير متطابقتين.')
            return render(request, 'core/change_password.html')

        try:
            validate_password(new_password, user=request.user)
        except ValidationError as e:
            for msg in e.messages:
                messages.error(request, msg)
            return render(request, 'core/change_password.html')

        if old_password == new_password:
            messages.error(request, 'كلمة المرور الجديدة يجب أن تكون مختلفة عن الحالية.')
            return render(request, 'core/change_password.html')

        request.user.set_password(new_password)
        request.user.save()

        auth_login(request, request.user)

        messages.success(request, 'تم تغيير كلمة المرور بنجاح.')
        return redirect('dashboard')

    return render(request, 'core/change_password.html')